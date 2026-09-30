"""
Generator script to compile the updated, comprehensive technical specification report
for EnviroMap-KG Objective 1 incorporating the latest production upgrades (Apache Jena Fuseki TDB2,
Celery/Redis task distribution, React 18/Vite WebSocket dashboard, strict ThingsBoard exclusion filter,
TimescaleDB integration, and LangGraph SQLite checkpointing).
"""

import os
import sys
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def add_callout(doc, text, title="IMPORTANT"):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.columns[0].width = Inches(6.5)
    cell = table.cell(0, 0)
    set_cell_background(cell, "F0F4F8")
    set_cell_margins(cell, top=140, bottom=140, left=200, right=200)
    
    # Left border highlight
    tcPr = cell._element.get_or_add_tcPr()
    borders = parse_xml(
        f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="none"/>'
        f'<w:left w:val="single" w:sz="24" w:space="0" w:color="0052CC"/>'
        f'<w:bottom w:val="none"/>'
        f'<w:right w:val="none"/>'
        f'</w:tcBorders>'
    )
    tcPr.append(borders)
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    run_title = p.add_run(f"[{title}] ")
    run_title.bold = True
    run_title.font.color.rgb = RGBColor(0, 82, 204)
    run_title.font.size = Pt(10)
    
    run_text = p.add_run(text)
    run_text.font.size = Pt(10)
    run_text.font.color.rgb = RGBColor(30, 41, 59)
    doc.add_paragraph()

def format_table(table, col_widths, headers, data):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    
    # Header row
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_background(hdr_cells[i], "1E293B")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=140, right=140)
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.bold = True
            run.font.color.rgb = RGBColor(255, 255, 255)
            run.font.size = Pt(9.5)
            run.font.name = "Arial"
            
    # Data rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.add_row().cells
        bg_color = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_background(row_cells[c_idx], bg_color)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=120, right=120)
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(30, 41, 59)
                run.font.name = "Arial"
                
    # Set widths
    for row in table.rows:
        for c_idx, width in enumerate(col_widths):
            row.cells[c_idx].width = Inches(width)

def build_document(output_paths):
    doc = Document()
    
    # Margins
    for section in doc.sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)
        
    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Calibri'
    normal_style.font.size = Pt(11)
    normal_style.font.color.rgb = RGBColor(30, 41, 59)
    normal_style.paragraph_format.line_spacing = 1.2
    normal_style.paragraph_format.space_after = Pt(6)

    # ── TITLE & HEADER ───────────────────────────────────────────────────────
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(16)
    p_title.paragraph_format.space_after = Pt(4)
    run_title = p_title.add_run("EnviroMap-KG: Standards-Aware Multi-Agent Semantic Interoperability Architecture")
    run_title.bold = True
    run_title.font.name = "Arial"
    run_title.font.size = Pt(21)
    run_title.font.color.rgb = RGBColor(15, 23, 42)

    p_sub = doc.add_paragraph()
    p_sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_sub.paragraph_format.space_after = Pt(14)
    run_sub = p_sub.add_run("Comprehensive Technical Specification & Study Report — Objective 1\nUpgraded Production Architecture with Apache Jena Fuseki TDB2, Celery Task Queues, & React WebSocket Interface")
    run_sub.font.size = Pt(12)
    run_sub.font.italic = True
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    p_meta = doc.add_paragraph()
    p_meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_meta.paragraph_format.space_after = Pt(20)
    run_meta = p_meta.add_run("Author / Project: EnviroMap-KG Engineering & Research  |  Document Version: 3.5 (Production Upgrade Edition)  |  Date: September 2026")
    run_meta.font.size = Pt(9.5)
    run_meta.font.bold = True
    run_meta.font.color.rgb = RGBColor(100, 116, 139)

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── 1. EXECUTIVE SUMMARY & ABSTRACT ──────────────────────────────────────
    doc.add_heading("1. Executive Summary & Abstract", level=1)
    doc.add_paragraph(
        "The proliferation of Internet of Things (IoT) monitoring platforms in smart agriculture, atmospheric monitoring, "
        "and industrial hydrology has produced acute semantic fragmentation. While edge microcontrollers accurately collect physical "
        "measurements, the telemetry remains trapped in isolated database silos. Disparate systems employ incompatible transport "
        "protocols (oneM2M, OPC UA, ThingsBoard, MQTT, REST), proprietary schema naming conventions ('temp', 'temperature', 'degC_val'), "
        "and conflicting domain assumptions (e.g., an unadorned 'temperature' column denotes liquid nutrient solution in hydroponics, "
        "but ambient air temperature in meteorological monitoring)."
    )
    doc.add_paragraph(
        "EnviroMap-KG Objective 1 addresses this challenge by establishing an autonomous, standards-aware, multi-agent semantic "
        "interoperability platform. The system ingests disparate PostgreSQL relational tables, introspects schema metadata, executes "
        "hybrid alignment against formal W3C (SSN/SOSA) and ETSI (SAREF, SAREF4WATR, SAREF4ENVI, SAREF4AGRI) ontologies, conducts "
        "multi-criteria reasoning, validates candidates against ontological constraints, reinforces successful alignments via episodic "
        "memory, translates raw measurements into canonical RDF triples, and persists them into a persistent, disk-backed Knowledge Graph."
    )
    doc.add_paragraph(
        "Current Upgrade Highlights:\n"
        "• Enterprise Triplestore: Migrated from in-memory RDFLib serialization to Apache Jena Fuseki with a TDB2 B-tree storage engine, "
        "scaling to billions of triples without RAM exhaustion.\n"
        "• Distributed Asynchronous Execution: Integrated Celery task workers backed by a Redis message broker for non-blocking per-domain pipeline execution.\n"
        "• Enterprise React Dashboard: Deployed a modern React 18 + Vite 5 + Recharts interface with live WebSocket telemetry push (/api/ws/telemetry) "
        "and dark/light mode persistence, replacing legacy Streamlit UI.\n"
        "• Strict Simulation Exclusion: Implemented multi-tier keyword filtering to eliminate 23+ million non-environmental ThingsBoard and "
        "Digital Twin rows, preserving strictly the authentic environmental testbeds.\n"
        "• Active Knowledge Graph Scale: Over 939,633 triples actively maintained with 100% validation accuracy (21/21 production fields)."
    )

    add_callout(
        doc,
        "EnviroMap-KG Objective 1 provides the authoritative, standards-compliant semantic foundation. Disparate sensor streams are unified into "
        "W3C SOSA / ETSI SAREF canonical observations stored in Apache Jena Fuseki TDB2, serving as the verified semantic input for downstream "
        "applications and Objective 2 autonomous cross-domain control.",
        title="OBJECTIVE 1 MISSION & UPGRADE STATUS"
    )

    # ── 2. LIVE PRODUCTION TESTBEDS & STRICT EXCLUSION FILTER ──────────────────
    doc.add_heading("2. Live Production Testbeds & Strict Platform Isolation", level=1)
    doc.add_paragraph(
        "On the production server (Ubuntu 24.04 LTS at 192.168.0.21), the PostgreSQL instance hosts diverse databases. A critical "
        "discovery during system scaling revealed that the server hosted internal simulation and test platforms alongside real environmental sensors:"
    )

    # Testbeds Table
    testbeds_data = [
        ["hydroponics", "onem2m_data", "oneM2M Edge Probes", "Hydroponics", "11,605+", "s4watr:Turbidity, s4watr:TotalDissolvedSolids, s4watr:HydrogenIonConcentration", "APPROVED (Live Telemetry)"],
        ["IAQ-5", "opcua_data", "OPC UA Industrial", "Air Quality", "41+", "s4envi:ParticulateMatter, s4envi:AirQualityIndex, saref:Temperature", "APPROVED (Live Telemetry)"],
        ["wq_bus_sensor", "thingsboard_data", "Water Bus Sensors", "Smart Water", "165+", "s4watr:Conductivity, s4watr:DissolvedOxygen, s4watr:WaterFlow", "APPROVED (Live Telemetry)"],
        ["thingsboard", "ts_kv / internal", "ThingsBoard Internal", "Simulation", "22,957,485", "N/A — Internal message queues and non-environmental telemetry", "EXCLUDED (Strict Filter)"],
        ["thingsboard_sim", "synthetic_data", "Synthetic Simulator", "Simulation", "937,076", "N/A — Synthetic benchmarking telemetry", "EXCLUDED (Strict Filter)"],
        ["digital_twin_dev_11", "dev_telemetry", "Virtual Asset Shadow", "Simulation", "266", "N/A — Digital Twin virtual state model", "EXCLUDED (Strict Filter)"]
    ]
    t = doc.add_table(rows=1, cols=7)
    format_table(t, [1.0, 0.9, 0.9, 0.8, 0.6, 1.3, 1.0], ["Database", "Table", "Source Protocol", "Domain", "Rows", "Canonical Target Concepts", "Pipeline Status"], testbeds_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    doc.add_paragraph(
        "Multi-Tier Exclusion Implementation:\n"
        "To prevent 23+ million non-environmental records from flooding the Knowledge Graph, strict exclusion rules were implemented across all layers:\n"
        "1. Database Discovery Level (api/main.py): Queries pg_database with EXCLUDED_KEYWORDS = ('postgres', 'camera', 'template', 'thingsboard', "
        "'thingsboard_sim', 'tb_', 'digital_twin', 'digitaltwin', 'sim'), immediately skipping unwanted databases.\n"
        "2. Domain Inference Level (infer_domain): Any candidate database or table token containing simulation markers returns None, preventing ingestion.\n"
        "3. Celery Worker Task Guard (orchestrator/tasks.py): Early-exit guard validates domain tokens before agent execution. If flagged, the worker returns "
        "status: 'skipped' without touching the triplestore.\n"
        "4. Frontend Presentation Level (full_app.jsx): Discovered Environments bar and TopKPI cards strictly filter domain totals to display only active environmental streams."
    )

    # ── 3. FULL TOOL STACK & 8-LAYER AGENTIC ARCHITECTURE ─────────────────────
    doc.add_heading("3. Full Tool Stack & 8-Layer Agentic AI Ecosystem", level=1)
    doc.add_paragraph(
        "The complete EnviroMap-KG platform is organized systematically into the 8-layer Open-Source Agentic AI hierarchy:"
    )

    layers_data = [
        ["Layer 1: Application", "React 18 + Vite 5 + Recharts + Tailwind CSS", "Modern reactive web dashboard with live WebSocket streaming (/api/ws/telemetry), dark/light theme persistence, interactive time-series plots, and SPARQL console."],
        ["Layer 2: API Gateway", "Python ThreadingHTTPServer + NGINX Proxy", "Dedicated API running on port 8008 (reverse-proxied via NGINX); provides /api/kg-data, /api/ws/telemetry, and /api/sparql with 5-second in-memory caching."],
        ["Layer 3: Orchestration", "LangGraph StateGraph + Celery / Redis", "StateGraph workflow with SQLite Checkpointing (data/pipeline_checkpoint.db) for crash recovery, coupled with Celery for parallel async domain tasks."],
        ["Layer 4: Intelligence", "HuggingFace MiniLM + LLM Advisory", "SentenceTransformer('all-MiniLM-L6-v2') 384-dim dense embeddings + Google Gemini 2.0 / Ollama for zero-shot semantic disambiguation when confidence < 0.75."],
        ["Layer 5: Protocol & DB", "PostgreSQL 16 + TimescaleDB + psycopg2", "PostgreSQL RDBMS with TimescaleDB hypertables for time-series partitioning; psycopg2 with RealDictCursor for metadata introspection."],
        ["Layer 6: Knowledge Store", "YAML Registry + JSON Episodic Memory", "35+ curated concepts in YAML, RapidFuzz lexical matcher, and data/mapping_memory.json for sub-millisecond fast-path mapping reuse."],
        ["Layer 7: Automation", "Linux systemd Services", "enviromap-pipeline.service running autonomous_worker.py every 360s with Restart=always and user linger enabled."],
        ["Layer 8: Triplestore", "Apache Jena Fuseki TDB2 + RDFLib", "Persistent disk-backed B-tree triplestore hosting 939,633+ triples with SPARQL 1.1 Update/Query HTTP endpoints; fallback to local Turtle (.ttl)."]
    ]
    t = doc.add_table(rows=1, cols=3)
    format_table(t, [1.4, 1.8, 3.3], ["Agentic Layer", "Technology Stack", "Role in EnviroMap-KG Upgraded Architecture"], layers_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── 4. DEEP-DIVE: THE SIX CORE AGENTS ─────────────────────────────────────
    doc.add_heading("4. Comprehensive Technical Breakdown of the Six Core Agents", level=1)
    doc.add_paragraph(
        "All agents inherit from BaseAgent (agents/base_agent.py), which enforces a strict execution contract: typed data models, "
        "microsecond performance timing, telemetry accumulation (execution count, total time, average latency), and structured logging."
    )

    # Agent 1
    doc.add_heading("4.1 Agent 1: Discovery / Profiler Agent (DiscoveryAgent)", level=2)
    doc.add_paragraph(
        "• Source Location: agents/discovery_agent.py | Execution Latency: ~96.9 ms\n"
        "• Primary Responsibility: Database introspection, statistical profiling, and environmental context extraction.\n"
        "• Input: Database name, schema, table name, source platform tag, and DataAcquisition instance.\n"
        "• Output: List[DiscoveryProfile] data contracts."
    )
    doc.add_paragraph(
        "Algorithmic Workflow:\n"
        "1. Schema Cataloging: Queries information_schema.columns to extract column names, SQL datatypes, and nullability.\n"
        "2. Column Filtering: Excludes primary keys ('id', 'uuid', 'pk') and timestamp columns, noting timestamp existence to verify streaming observability.\n"
        "3. Statistical Ingestion: Computes min, max, null count, and distinct value cardinality over sample rows.\n"
        "4. Sibling Context Synthesis: Identifies all neighbouring measurement columns in the same table. For instance, discovering 'tds' and 'ph' "
        "alongside 'temperature' provides decisive contextual proof that 'temperature' belongs to the aqueous hydroponic domain rather than meteorology."
    )

    # Agent 2
    doc.add_heading("4.2 Agent 2: Alignment Agent (AlignmentAgent)", level=2)
    doc.add_paragraph(
        "• Source Location: agents/alignment_agent.py | Execution Latency: ~181.5 ms\n"
        "• Primary Responsibility: Multi-modal lexical and dense semantic vector alignment against the Ontology Registry.\n"
        "• Input: DiscoveryProfile, OntologyRegistry reference, EmbeddingModel reference.\n"
        "• Output: List[AlignmentCandidate] ranked descending by combined similarity score."
    )
    doc.add_paragraph(
        "Mathematical Model & Scoring Formulation:\n"
        "Given a source field profile P with field name f, domain D, and neighbouring fields N = {n_1, ..., n_k}:\n"
        "1. Dense Semantic Vector Embedding (S_semantic):\n"
        "   Using sentence-transformers/all-MiniLM-L6-v2 (384-dimensional dense representation), query vector q is generated by concatenating "
        "the field name, domain, and neighbour tokens: q = Embed(f || D || N). Cosine similarity against precomputed concept vector c_C is computed:\n"
        "   S_semantic(f, C) = (q · c_C) / (||q||_2 * ||c_C||_2)\n"
        "2. Lexical Token Similarity (S_lexical):\n"
        "   Computed via RapidFuzz Levenshtein Token-Sort Ratio against the concept label and curated aliases:\n"
        "   S_lexical(f, C) = max_{a in Aliases(C)} [ TokenSortRatio(f, a) / 100 ]\n"
        "3. Domain Applicability Boost (S_domain):\n"
        "   S_domain(D, C) = 1.0 if D in DomainApplicability(C) else 0.0\n"
        "4. Combined Alignment Function (S_combined):\n"
        "   S_combined = (0.35 * S_lexical) + (0.45 * S_semantic) + (0.20 * S_domain)\n"
        "5. Exact-Match Shortcut: If f matches an alias exactly (normalized lowercase), S_combined = 1.0 with match_method = 'exact'."
    )

    # Agent 3
    doc.add_heading("4.3 Agent 3: Mapping Agent (MappingAgent)", level=2)
    doc.add_paragraph(
        "• Source Location: agents/mapping_agent.py | Execution Latency: < 0.05 ms\n"
        "• Primary Responsibility: Evidence aggregation and context packaging.\n"
        "• Input: DiscoveryProfile, List[AlignmentCandidate].\n"
        "• Output: Immutable MappingRecord encapsulating top-K candidates, value ranges, source table paths, and neighbouring environmental markers."
    )

    # Agent 4
    doc.add_heading("4.4 Agent 4: Reasoning Agent (ReasoningAgent)", level=2)
    doc.add_paragraph(
        "• Source Location: agents/reasoning_agent.py | Execution Latency: ~2,411.2 ms (includes optional LLM call)\n"
        "• Primary Responsibility: Multi-Criteria Decision Analysis (MCDA) with zero-shot LLM Advisory fallback.\n"
        "• Input: MappingRecord, OntologyRegistry, LLMAdvisory instance.\n"
        "• Output: ReasoningDecision containing selected concept URI, confidence score, criteria evidence, and explanation chain."
    )
    doc.add_paragraph(
        "Multi-Criteria Scoring Formula:\n"
        "For candidate C_i, confidence score K(C_i) is evaluated:\n"
        "   K(C_i) = 0.30 * S_combined(C_i) + Delta_domain + Delta_unit + Delta_type + Delta_context + Delta_range\n"
        "Where bonuses are granted:\n"
        "   • Delta_domain (+0.20): Target concept explicitly supports the source domain.\n"
        "   • Delta_unit (+0.15): Inferred or declared unit matches concept compatible units.\n"
        "   • Delta_type (+0.10): Database column type matches allowed ontological datatypes.\n"
        "   • Delta_context (+0.15): Co-occurring fields reinforce domain context (e.g., 'ph' confirming water).\n"
        "   • Delta_range (+0.10): Sample readings fall strictly within physical boundaries [val_min, val_max].\n"
        "\nLLM Advisory Fallback Trigger:\n"
        "If max_i K(C_i) < 0.75, the agent triggers components/llm_advisory.py. A structured JSON payload containing field metadata, "
        "sample readings, candidate URIs, and neighbour tokens is submitted to Google Gemini 2.0 / Ollama. The model returns structured "
        "JSON with recommended concept URI and rationalization chain, which is recorded in the audit log."
    )

    # Agent 5
    doc.add_heading("4.5 Agent 5: Validation Agent (ValidationAgent)", level=2)
    doc.add_paragraph(
        "• Source Location: agents/validation_agent.py | Execution Latency: < 0.05 ms\n"
        "• Primary Responsibility: Authoritative, zero-tolerance constraint gatekeeper.\n"
        "• Input: ReasoningDecision, MappingRecord, OntologyRegistry.\n"
        "• Output: ValidationResult (VALID, INVALID, CONDITIONAL)."
    )

    val_rules = [
        ["ontology_consistency", "Required (Critical)", "Target concept URI must exist in Ontology Registry", "Mapping immediately REJECTED (INVALID)"],
        ["domain_consistency", "Required (Critical)", "Source domain must match concept DomainApplicability", "Mapping immediately REJECTED (INVALID)"],
        ["datatype_compatibility", "Required (Critical)", "PostgreSQL datatype must match concept compatible_datatypes", "Mapping immediately REJECTED (INVALID)"],
        ["unit_compatibility", "Optional (Heuristic)", "Unit must match concept compatible_units", "Confidence penalized by -0.10"],
        ["range_plausibility", "Optional (Heuristic)", "Sample readings must fall within [val_min, val_max]", "Sensor offline markers (-1.0/-127.0) tagged, penalty -0.10"],
        ["contextual_consistency", "Optional (Heuristic)", "Neighbouring tokens must not contradict property definitions", "Confidence penalized by -0.05"]
    ]
    t = doc.add_table(rows=1, cols=4)
    format_table(t, [1.5, 1.4, 2.1, 1.5], ["Validation Check", "Constraint Level", "Condition Evaluated", "Enforcement Action"], val_rules)
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    # Agent 6
    doc.add_heading("4.6 Agent 6: Learning / Memory Agent (LearningAgent)", level=2)
    doc.add_paragraph(
        "• Source Location: agents/learning_agent.py | Execution Latency: ~13.5 ms\n"
        "• Primary Responsibility: Episodic memory persistence and deterministic fast-path acceleration.\n"
        "• Input: ValidationResult, ReasoningDecision, DiscoveryProfile.\n"
        "• Storage: data/mapping_memory.json"
    )
    doc.add_paragraph(
        "Reinforcement & Exponential Moving Average (EMA) Formulation:\n"
        "When a mapping achieves validated status with confidence >= 0.85:\n"
        "1. Cache Check: If field key 'domain.table.field' exists, the confidence is updated via EMA:\n"
        "   C_updated = 0.20 * C_current + 0.80 * C_stored\n"
        "   use_count = use_count + 1\n"
        "   last_used = ISO_8601_TIMESTAMP\n"
        "2. Deterministic Reuse: In subsequent execution cycles, RuleBasedEngine checks mapping_memory.json during intake. "
        "Confirmed mappings are resolved in < 0.08 ms with 100% confidence, completely bypassing embedding matrix multiplications "
        "and LLM inference overhead."
    )

    # ── 5. ENTERPRISE UPGRADE IMPLEMENTATIONS ─────────────────────────────────
    doc.add_heading("5. Upgraded Production Infrastructure & Components", level=1)

    doc.add_heading("5.1 Apache Jena Fuseki with TDB2 Persistent Storage", level=2)
    doc.add_paragraph(
        "To transition from prototype in-memory RDFLib storage to an enterprise-grade triplestore capable of hosting millions of triples:\n"
        "• Dedicated Triplestore Daemon: Apache Jena Fuseki runs as a native system service on port 3030 with a disk-backed TDB2 dataset (/enviromap).\n"
        "• Jena Connector (knowledge_graph/jena_connector.py): Encapsulates HTTP SPARQL Update (INSERT DATA) and SPARQL Query protocols with retry logic.\n"
        "• Persistent B-tree Indexing: Retains graph structure across server reboots, scaling seamlessly to 10+ million triples without memory saturation.\n"
        "• Transparent Fallback: If Fuseki is unreachable, the system automatically writes triples to local Turtle files (data/output/enviromap_kg.ttl)."
    )

    doc.add_heading("5.2 Distributed Asynchronous Task Queue: Celery & Redis", level=2)
    doc.add_paragraph(
        "• Celery Application (orchestrator/celery_app.py): Configured with a Redis message broker (redis://127.0.0.1:6379/0) and result backend.\n"
        "• Parallel Domain Tasks (orchestrator/tasks.py): Each discovered environmental database (hydroponics, IAQ-5, wq_bus_sensor) is dispatched "
        "as an independent asynchronous Celery task (async_run_domain_pipeline.delay()), executing the 6-agent pipeline in parallel across CPU cores.\n"
        "• Resilient Fallback: If Redis is offline, tasks gracefully degrade to non-blocking Python daemon threads without halting the system."
    )

    doc.add_heading("5.3 Modern Enterprise Frontend: React 18 + Vite 5 Dashboard", level=2)
    doc.add_paragraph(
        "• Architecture: Full React 18 single-page application built with Vite 5, packaged into react_dist.tar.gz and served via NGINX on port 80.\n"
        "• Real-Time Telemetry Streaming: Connects via native WebSocket to /api/ws/telemetry, receiving sensor deltas every 5 seconds without page reloading.\n"
        "• Interactive Analytics: Renders dynamic time-series charts using Recharts, dark/light theme switching with localStorage persistence, "
        "and an integrated SPARQL query explorer connected to Apache Jena Fuseki."
    )

    doc.add_heading("5.4 Dedicated Backend API Gateway (api/main.py)", level=2)
    doc.add_paragraph(
        "Runs on port 8008 under ThreadingHTTPServer. Exposes REST and WebSocket interfaces:\n"
        "• GET /api/kg-data: Returns cached aggregated telemetry, domain totals, and mapping provenance (with 5-second TTL cache).\n"
        "• WS /api/ws/telemetry: RFC 6455 compliant WebSocket stream broadcasting live sensor updates.\n"
        "• GET /api/domains: Returns active environmental domains with simulation platforms filtered out.\n"
        "• POST /api/sparql: Proxies SPARQL queries directly to Apache Jena Fuseki, returning standard W3C SPARQL JSON results."
    )

    doc.add_heading("5.5 LangGraph StateGraph with SQLite Checkpointing", level=2)
    doc.add_paragraph(
        "The workflow orchestration utilizes LangGraph's StateGraph coupled with SqliteSaver (persisting to data/pipeline_checkpoint.db). "
        "Every domain processing phase commits state snapshots to SQLite. In the event of power loss or server reboots, the pipeline "
        "resumes execution from the exact interrupted node, eliminating redundant re-processing."
    )

    # ── 6. PRODUCTION 24/7 OPERATION & VERIFICATION ──────────────────────────
    doc.add_heading("6. Autonomous 24/7 Production Operation & Benchmarks", level=1)
    doc.add_paragraph(
        "On the Ubuntu 24.04 server (192.168.0.21), the platform is governed by dual systemd daemons and user linger:\n"
        "1. enviromap-pipeline.service: Executes scripts/autonomous_worker.py every 360 seconds (6 minutes), advancing watermarks in data/watermark.json.\n"
        "2. enviromap-web.service: Manages the API gateway and dashboard backend on port 8008.\n"
        "3. Ubuntu Linger: Enabled via sudo loginctl enable-linger cse-dept, preventing process termination upon SSH logout."
    )

    benchmarks_data = [
        ["Discovery Agent", "96.9 ms", "Automated PostgreSQL information_schema profiling + 100 sample rows."],
        ["Alignment Agent", "181.5 ms", "384-dim tensor dot-product + RapidFuzz Levenshtein scoring."],
        ["Mapping Agent", "< 0.05 ms", "In-memory candidate packaging & context dictionary assembly."],
        ["Reasoning Agent", "2,411.2 ms", "Multi-criteria scoring matrix evaluation & optional LLM advisory call."],
        ["Validation Agent", "< 0.05 ms", "Execution of 6-point constraint suite across all mapped columns."],
        ["Learning Agent", "13.5 ms", "EMA confidence recalculation & atomic JSON disk persistence."],
        ["Fast-Path Memory Lookup", "< 0.08 ms", "100% deterministic hash lookup for known approved mappings."],
        ["Total Pipeline Runtime", "67.5 s", "End-to-end processing across live databases (583 rows, 4,468 observations)."],
        ["Validation Success Rate", "21 / 21 (100%)", "Zero invalid mappings across production tables."],
        ["Total Triplestore Size", "939,633 Triples", "Active SOSA/SAREF knowledge graph in Apache Jena Fuseki TDB2."]
    ]
    t = doc.add_table(rows=1, cols=3)
    format_table(t, [1.6, 1.4, 3.5], ["Operational Stage / Metric", "Measured Benchmark", "Technical Significance & Notes"], benchmarks_data)
    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # ── 7. SUMMARY & ACADEMIC VALUE ───────────────────────────────────────────
    doc.add_heading("7. Summary & Academic Conclusions", level=1)
    doc.add_paragraph(
        "EnviroMap-KG Objective 1 demonstrates that heterogeneous environmental IoT streams can be semantically unified without "
        "manual schema engineering and without unconstrained LLM hallucinations. By coupling:\n"
        "1. Statistical schema profiling (DiscoveryAgent),\n"
        "2. Dense vector embeddings with domain boosting (AlignmentAgent),\n"
        "3. Multi-criteria evidence evaluation (ReasoningAgent),\n"
        "4. Strict constraint validation (ValidationAgent),\n"
        "5. Episodic memory reinforcement (LearningAgent),\n"
        "6. Crash-resilient graph orchestration (LangGraph StateGraph),\n"
        "7. Scalable disk-backed triplestores (Apache Jena Fuseki TDB2), and\n"
        "8. Real-time distributed interfaces (React 18 + WebSocket + Celery/Redis),\n\n"
        "the platform establishes an autonomous, self-improving semantic bridge between raw IoT sensor databases and standardized W3C SOSA / "
        "ETSI SAREF Knowledge Graphs, fulfilling all requirements of Objective 1 and providing the validated semantic stream for Objective 2."
    )

    # Save to all target paths with fallback for file locks
    for path in output_paths:
        os.makedirs(os.path.dirname(path), exist_ok=True)
        try:
            doc.save(path)
            print(f"Document saved successfully to: {path}")
        except PermissionError:
            # File is currently open in Microsoft Word
            alt_path = path.replace(".docx", "_v2.docx")
            doc.save(alt_path)
            print(f"Original file was open in Word. Saved to alternative path: {alt_path}")

if __name__ == "__main__":
    download_dir = r"C:\Users\Shreenithi\Downloads"
    artifact_dir = r"C:\Users\Shreenithi\.gemini\antigravity\brain\09c39602-6cb4-4017-9b53-06e87a8af5a5"
    
    doc_paths = [
        os.path.join(download_dir, "EnviroMap_KG_Objective_1_Upgraded_Technical_Report.docx"),
        os.path.join(download_dir, "EnviroMap_KG_Objective_1_Complete_Technical_Report.docx"),
        os.path.join(artifact_dir, "EnviroMap_KG_Objective_1_Complete_Technical_Report.docx")
    ]
    build_document(doc_paths)
